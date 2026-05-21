"""高德 POI 类型编码 → 中文名转换。

从 Amap_poicode.xlsx 加载映射，将 "060102" 转为 "购物服务;商场;普通商场"。"""

import openpyxl
from pathlib import Path

_CODE_XLSX = Path(__file__).parent.parent.parent / "assert" / "Amap_poicode.xlsx"

_type_cache: dict[str, str] | None = None


def _load_codes() -> dict[str, str]:
    global _type_cache
    if _type_cache is not None:
        return _type_cache

    if not _CODE_XLSX.exists():
        raise FileNotFoundError(f"POI 编码表不存在: {_CODE_XLSX}")

    try:
        wb = openpyxl.load_workbook(_CODE_XLSX, read_only=True)
    except Exception:
        raise RuntimeError(f"POI 编码表读取失败，请检查文件是否损坏: {_CODE_XLSX}")
    ws = wb.active
    mapping: dict[str, str] = {}

    for r in range(2, ws.max_row + 1):
        new_type = ws.cell(r, 2).value
        if new_type is None:
            continue
        new_type = str(new_type)
        big = ws.cell(r, 3).value or ""
        mid = ws.cell(r, 4).value or ""
        sub = ws.cell(r, 5).value or ""
        code = new_type
        if big and mid and sub:
            mapping[code] = f"{big};{mid};{sub}"
        elif big and mid:
            mapping[code] = f"{big};{mid}"
        elif big:
            mapping[code] = big

    wb.close()
    _type_cache = mapping
    return mapping


def code_to_name(code: str) -> str:
    """将 POI 编码转为 "大类;中类;小类" 格式。"""
    if not code:
        return ""
    codes = _load_codes()
    return codes.get(code, code)


if __name__ == "__main__":
    assert code_to_name("060102") == "购物服务;商场;普通商场"
    assert code_to_name("140600") == "科教文化服务;科技馆;科技馆"
    assert code_to_name("110210") == "风景名胜;风景名胜;红色景区"
    assert code_to_name("999999") == "999999"  # 查不到返回原码
    assert code_to_name("") == ""
    print("all pass")
